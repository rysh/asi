"""Properties of the file that actually ships.

Reading the rendered deck back catches what the content tests cannot: fonts
falling back to a theme default, sizes left to inheritance, geometry lost in
translation.
"""

import pytest

pytest.importorskip("pptx")

from pptx.util import Emu  # noqa: E402

from slides.content import SLIDES  # noqa: E402
from slides.render import render_deck  # noqa: E402
from slides.theme import SLIDE_H, SLIDE_W, THEME  # noqa: E402


@pytest.fixture(scope="module")
def deck():
    return render_deck(SLIDES)


def runs(deck):
    for slide in deck.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    yield run


def test_deck_is_sixteen_by_nine(deck):
    assert deck.slide_width == Emu(SLIDE_W)
    assert deck.slide_height == Emu(SLIDE_H)


def test_deck_has_every_slide(deck):
    assert len(deck.slides) == len(SLIDES) == 29


def test_every_run_names_its_font(deck):
    """An unset font name means the PowerPoint theme decides — which is what
    breaks when the deck is opened in Keynote or Google Slides."""
    for run in runs(deck):
        assert run.font.name == THEME.font_latin, (run.font.name, repr(run.text[:40]))


def test_every_run_sets_its_size(deck):
    for run in runs(deck):
        assert run.font.size is not None, repr(run.text[:40])


def test_no_run_falls_below_the_font_floor(deck):
    """Footnotes, page numbers and caveats are allowed to be smaller."""
    small = {THEME.footnote_pt, THEME.caveat_pt, THEME.page_pt}
    for run in runs(deck):
        pt = run.font.size.pt
        assert pt >= THEME.min_body_pt or pt in small, (pt, run.text[:40])


def test_four_conditions_diagram_is_drawn(deck):
    index = next(
        i for i, s in enumerate(SLIDES) if "Four Conditions" in s.title
    )
    slide = deck.slides[index]
    text = " ".join(
        run.text
        for shape in slide.shapes
        if shape.has_text_frame
        for para in shape.text_frame.paragraphs
        for run in para.runs
    )
    assert text.count("✓") == 2
    assert text.count("✗") == 2
    for name in ("SPEED", "COLLECTIVE INTELLIGENCE",
                 "GENERAL WISDOM UNDER COMPLEXITY", "SCIENTIFIC CREATIVITY"):
        assert name in text


def test_deck_carries_its_five_figures(deck):
    """The five figures are embedded, not linked — a linked picture arrives at
    the venue as an empty frame."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    pictures = [
        shape
        for slide in deck.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert len(pictures) == 5
    for pic in pictures:
        assert pic.image.blob, "figure is linked rather than embedded"


def test_deck_stays_within_size_budget(deck, tmp_path):
    """Five full-bleed PNGs dominate the file; anything far above that means
    something unintended got in."""
    out = tmp_path / "deck.pptx"
    deck.save(out)
    assert out.stat().st_size < 12_000_000
